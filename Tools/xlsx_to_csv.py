# -*- coding: utf-8 -*-
"""기획표.xlsx의 시트를 Tools/balance/*.csv로 찍어낸다 [지시 2026-10-13].

**무엇을 위한 것인가.** 기획자가 값을 고치는 곳은 엑셀이고, 유니티가 읽는 것은 CSV다.
사이를 손으로 옮기면 어느 쪽이 최신인지 금방 모르게 된다. 시트 하나가 CSV 하나다.

**주석을 지키는 것이 이 도구의 핵심이다.** CSV 맨 위의 `#` 블록에는 실측 결과와
그 값을 그렇게 정한 이유가 적혀 있다 — 엑셀에는 그게 없다. 통째로 덮어쓰면
프로젝트에서 제일 비싼 기록이 매번 날아간다. 그래서 **주석은 그대로 두고 표만 바꾼다.**

**헤더가 다르면 멈춘다.** 엑셀이 오래됐는데 CSV만 바뀐 경우가 실제로 있다
(2026-10-14의 grades·score_rules). 열 이름이 어긋나면 그 시트는 건너뛰고 알린다 —
조용히 덮어쓰면 오늘 한 작업이 어제 값으로 되돌아간다.

**기본은 미리보기다.** 쓰려면 --write를 붙여야 한다.

쓰는 법:
    python Tools/xlsx_to_csv.py              # 뭐가 달라지는지만 보여 준다
    python Tools/xlsx_to_csv.py --write      # 실제로 쓴다
    python Tools/xlsx_to_csv.py --write structures zone    # 그 시트만
"""
import io
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
BOOK = os.path.join(HERE, 'balance', '기획표.xlsx')
OUT = os.path.join(HERE, 'balance')

# 표가 아닌 시트 — 읽기만 하는 곳이다
SKIP = {'읽어주세요', '계산'}


def cell(v):
    """엑셀 값을 CSV 글자로. 정수로 떨어지는 실수는 정수로 적는다 — 2.0이 아니라 2다."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def read_sheet(ws):
    """첫 줄을 헤더로, 그 아래를 행으로 읽는다. 완전히 빈 줄은 버린다.

    **행은 헤더 길이에 맞춰 채운다.** 뒤쪽 빈 칸을 떼면 `a,b,c,,`가 `a,b,c`가 되어
    열이 밀린다 — zone.csv의 timeLimit/quota가 첫 줄에만 있어서 바로 깨졌다.
    """
    rows = []
    header = None
    for raw in ws.iter_rows(values_only=True):
        vals = [cell(v) for v in raw]
        trimmed = list(vals)
        while trimmed and trimmed[-1] == '':
            trimmed.pop()
        if not trimmed:
            continue

        if header is None:
            header = trimmed
            continue

        row = vals[:len(header)]
        while len(row) < len(header):
            row.append('')
        rows.append(row)
    return header, rows


def split_csv(path):
    """CSV를 (맨 위 주석 블록, 헤더, 나머지 줄)로 가른다."""
    if not os.path.exists(path):
        return [], None, []

    with io.open(path, encoding='utf-8-sig') as f:
        lines = f.read().splitlines()

    comments, header, rest = [], None, []
    for line in lines:
        if header is None:
            if line.startswith('#') or line.strip() == '':
                comments.append(line)
            else:
                header = line
        else:
            rest.append(line)
    return comments, header, rest


def tail_comments(rest):
    """표 뒤에 붙은 주석 줄. structures.csv의 '구역에 없는 종류' 블록 같은 것."""
    return [l for l in rest if l.startswith('#')]


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    write = '--write' in sys.argv
    only = [a for a in sys.argv[1:] if not a.startswith('--')]

    if not os.path.exists(BOOK):
        print('엑셀을 못 찾았다:', BOOK)
        return 1

    wb = load_workbook(BOOK, data_only=True)
    changed = skipped = untouched = 0

    for name in wb.sheetnames:
        if name in SKIP:
            continue
        if only and name not in only:
            continue

        path = os.path.join(OUT, name + '.csv')
        header, rows = read_sheet(wb[name])
        if header is None:
            print('[건너뜀] %-14s 시트가 비어 있다' % name)
            skipped += 1
            continue

        comments, old_header, old_rest = split_csv(path)

        if old_header is None:
            print('[새 파일] %-14s %d행' % (name, len(rows)))
        elif [h.strip() for h in old_header.split(',')] != header:
            print('[멈춤]   %-14s 열 이름이 다르다 — 덮어쓰지 않는다' % name)
            print('           CSV  : ' + old_header)
            print('           엑셀 : ' + ','.join(header))
            skipped += 1
            continue

        body = [','.join(r) for r in rows]
        keep = tail_comments(old_rest)
        out = comments + [','.join(header)] + body + keep

        old_body = [l for l in old_rest if not l.startswith('#')]
        if old_header is not None and old_body == body:
            print('[같음]   %-14s %d행' % (name, len(rows)))
            untouched += 1
            continue

        print('[바뀜]   %-14s %d행 (전 %d행)' % (name, len(rows), len(old_body)))
        for i, line in enumerate(body):
            if i >= len(old_body):
                print('           + ' + line[:70])
            elif old_body[i] != line:
                print('           - ' + old_body[i][:70])
                print('           + ' + line[:70])
        changed += 1

        if write:
            with io.open(path, 'w', encoding='utf-8-sig', newline='\r\n') as f:
                f.write('\n'.join(out) + '\n')

    print()
    print('바뀜 %d · 같음 %d · 건너뜀 %d' % (changed, untouched, skipped))
    if changed and not write:
        print('실제로 쓰려면 --write 를 붙일 것.')
    if write and changed:
        print('유니티에서 Demo/5로 불러와야 게임에 들어간다.')
    return 0


if __name__ == '__main__':
    sys.exit(main())
